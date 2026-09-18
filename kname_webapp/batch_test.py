#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
실제 API 키가 있는 환경에서 여러 이름을 한 번에 돌려본다.

    $env:ANTHROPIC_API_KEY="..."
    python batch_test.py                       # 기본 20개
    python batch_test.py --file names.txt      # 파일에서 읽기
    python batch_test.py --names "Siobhan Kowalski 여" "Tadhg Byrne 남"
    python batch_test.py --n 50                # 사전에서 무작위 50개

결과는 batch_result.txt 와 batch_result.json 에 저장된다.
JSON 을 대화에 붙여 넣으면 함께 검토할 수 있다.
"""
import argparse
import json
import os
import random
import re
import sys

BASE = os.path.dirname(os.path.abspath(__file__))
os.chdir(BASE)
sys.path.insert(0, BASE)

HANGUL = re.compile(r'[가-힣]')
HANJA = re.compile(r'[\u4e00-\u9fff]')
LATIN_EXT = re.compile(r'[\u00c0-\u024f]')

# 기본 세트 — 사전에 없을 법한 다양한 어원
DEFAULT = [
    ('Siobhan', 'Kowalski', '여'), ('Tadhg', 'Byrne', '남'),
    ('Aoibheann', 'Fitzpatrick', '여'), ('Padraig', 'Nakagawa', '남'),
    ('Oluwaseun', 'Adebayo', '남'), ('Chidinma', 'Nwosu', '여'),
    ('Bartosz', 'Lindqvist', '남'), ('Katarzyna', 'Wojcik', '여'),
    ('Thibault', 'Beaulieu', '남'), ('Clemence', 'Chevalier', '여'),
    ('Dimitrios', 'Papadopoulos', '남'), ('Svetlana', 'Kuznetsov', '여'),
    ('Torbjorn', 'Rasmussen', '남'), ('Solveig', 'Bjornstad', '여'),
    ('Meenakshi', 'Venkatesan', '여'), ('Harpreet', 'Bhullar', '남'),
    ('Mustafa', 'Alsaadi', '남'), ('Zahra', 'Haddad', '여'),
    ('Kenjiro', 'Hasegawa', '남'), ('Wilhelmina', 'Oosterhuis', '여'),
]


def _save(path, writer):
    """
    파일을 저장한다. 열려 있어 잠긴 경우(엑셀·메모장) 죽지 않고
    'result_brazil (2).xlsx' 처럼 옆 이름으로 저장하고 그 경로를 돌려준다.

    Windows 는 열려 있는 파일을 덮어쓰지 못한다(PermissionError). 테스트를
    돌릴 때마다 엑셀을 닫아야 하는 것은 번거롭고, 결과를 통째로 잃는 것은
    더 나쁘다.
    """
    stem, ext = os.path.splitext(path)
    for i in range(1, 20):
        target = path if i == 1 else f'{stem} ({i}){ext}'
        try:
            writer(target)
            if target != path:
                print(f'  ! {os.path.basename(path)} 이(가) 열려 있어 '
                      f'{os.path.basename(target)} 로 저장했습니다.')
            return target
        except PermissionError:
            continue
        except OSError as e:
            print(f'  ! {os.path.basename(target)} 저장 실패: {e}')
            return None
    print(f'  ! {os.path.basename(path)} 저장 실패 — 열려 있는 파일을 닫아주세요.')
    return None


def parse_line(s):
    """'Siobhan Kowalski 여' → ('Siobhan', 'Kowalski', '여')"""
    parts = s.strip().split()
    if len(parts) < 2:
        return None
    sex = '여'
    if parts[-1] in ('남', '여', 'other', 'either', 'M', 'F'):
        sex = {'M': '남', 'F': '여', 'either': 'other'}.get(parts[-1], parts[-1])
        parts = parts[:-1]
    if len(parts) < 2:
        return None
    return (parts[0], ' '.join(parts[1:]), sex)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--file', help='한 줄에 "이름 성 성별" 형식')
    ap.add_argument('--names', nargs='*', help='직접 지정')
    ap.add_argument('--n', type=int, help='사전에서 무작위 N개')
    ap.add_argument('--out', help='결과 파일 이름(확장자 제외). '
                                  '생략하면 --file 이름에서 자동으로 정한다')
    args = ap.parse_args()

    import app as A

    if not A.TRANSLIT.llm_available:
        print('=' * 74)
        print('⚠ ANTHROPIC_API_KEY 가 없습니다 — LLM 을 한 번도 호출하지 않습니다.')
        print('  의미 설명은 미리 작성된 602개 사전이나 로컬 템플릿에서만 채워지므로,')
        print('  이 상태의 결과로는 LLM 품질을 판단할 수 없습니다.')
        print('  $env:ANTHROPIC_API_KEY="sk-ant-..." 설정 후 다시 실행하세요.')
        print('=' * 74 + '\n')
    else:
        print('✔ ANTHROPIC_API_KEY 확인 — LLM 경로로 실행합니다.\n')

    # 의미 설명의 출처를 판정한다. 이게 없으면 "설명이 나왔다"는 사실만 알 뿐
    # 그게 LLM 이 쓴 것인지 사전에서 꺼낸 것인지 구분할 수 없다.
    # 확실한 판정 근거는 app.py 가 채워 주는 meaning_error 다. 아래 문구 목록은
    # 그것이 없을 때를 위한 보조 신호일 뿐이며, 넣으면 안 되는 것이 두 종류 있다.
    #   · 'It pictures someone'  — 602개 사전 설명 중 272개가 쓴다
    #   · _sound_note() 의 세 문구('Balanced in sound…', 'Firm and grounded…',
    #     'Soft and open…') — 순우리말 이름 설명(_native_desc)이 템플릿과
    #     같은 헬퍼를 공유해 정상 결과에도 그대로 나온다
    # 템플릿은 사전과 같은 문체로 쓰였으므로 문체만으로는 구분되지 않는다.
    _TPL = ('Together they picture', 'Together they bring',
            'Together they speak of', 'It speaks of one who',
            'The whole name reads calm and considered')

    def _norm(s):
        """성별 치환(boy↔girl, he↔she)을 무시하고 비교하기 위한 정규화."""
        s = re.sub(r'\s+', ' ', s or '').strip()
        return re.sub(r'\b(boys?|girls?|his|her|hers|he|she|him|sons?|'
                      r'daughters?|men|women)\b', '~', s, flags=re.I)

    # 순우리말 이름 설명(_native_desc)만 쓰는 문구. LLM 도 사전도 아닌
    # 로컬 생성이므로 따로 센다(602 사전·템플릿 어디에도 나오지 않는다).
    _NATIVE = ('the meaning lives right in the sound',            # _native_desc
               'Unlike Sino-Korean names, it is built from a Korean word')

    def meaning_source(given, text, err, declared=None):
        # ★ app.py 가 알려주는 출처가 있으면 그대로 쓴다(추측 없음).
        #   문체로 되짚는 아래 경로는 구버전 app.py 를 위한 보조일 뿐이다.
        if declared:
            return declared
        if not text:
            return 'none'
        if err or any(s in text for s in _TPL):
            return 'template'
        if any(x in text for x in _NATIVE):
            return 'native'
        # ② 미리 작성된 602개 원문과 일치하는가 (성별 치환분은 무시하고 비교)
        nt = _norm(text)
        for sx in ('male', 'female'):
            info = getattr(A, 'GIVEN_INFO', {}).get((sx, given))
            if not info:
                continue
            nd = _norm(info.get('meaning_en') or '')
            if nd and (nd == nt or nd[:150] == nt[:150]
                       or nt[:150] in nd or nd[:150] in nt):
                return 'dict'
        # ③ 둘 다 아니면 LLM 이 새로 쓴 것
        return 'llm'

    # 템플릿이 만들어내는 영어 문법 오류(관사·품사 오분류) 탐지
    BAD_EN = re.compile(
        r'\b(?:a|an) (?:poetry|true|wisdom|virtue|grace|jade|gold|silk|water|'
        r'sunlight|moonlight|robust|crimson|second|sky|snow|rain|earth)\b')

    # 독자를 'you/your' 로 부르는 문장. 생성 프롬프트와 검수 체크리스트가
    # 모두 금지하는데도 "Your name is Deacon." 이 한 번 통과한 적이 있다.
    SECOND_PERSON = re.compile(r'\b(?:you|your|yours|you\'re)\b', re.I)

    def char_form_misses(given, hanja, text):
        """글자별 표기 형식('혜 (惠, hye)')을 지키지 않은 글자 목록.

        형식은 프롬프트가 지시하고 검수가 확인하지만, 둘 다 LLM 이므로
        '지켜졌는지'는 코드로 세어야 안다. 문구 감사(tools/audit_reason_phrases.py)와
        같은 취지다 — LLM 이 지킬 규칙은 결정론적으로 검사한다.
        """
        from pronounce_guide import romanize_syllable
        miss = []
        for i, syl in enumerate(given or ''):
            if i >= len(hanja or ''):
                break
            hj = hanja[i]
            if not hj:
                continue
            rom = romanize_syllable(syl, capitalize=False)
            want = f'{syl} ({hj}, {rom})' if rom else f'{syl} ({hj})'
            if want not in (text or ''):
                miss.append(syl)
        return miss

    # 대상 결정
    if args.file:
        cases = [c for c in (parse_line(l) for l in open(args.file, encoding='utf-8')) if c]
    elif args.names:
        cases = [c for c in (parse_line(x) for x in args.names) if c]
    elif args.n:
        n2t = json.load(open('data/dict_name_to_translit.json'))
        random.seed()
        cases = []
        for sex, sk in (('male', '남'), ('female', '여')):
            for en in random.sample(list(n2t[sex]), args.n // 2):
                cases.append((en.title(), 'Smith', sk))
    else:
        cases = DEFAULT

    out_lines, rows = [], []

    def say(s=''):
        print(s)
        out_lines.append(s)

    say('=' * 74)
    say(f'배치 테스트 {len(cases)}건')
    say('=' * 74)

    warn = 0
    for i, (first, last, sex) in enumerate(cases, 1):
        d = A.convert_name(first, last, sex)
        if 'error' in d:
            say(f'{i:3d}. ❌ {first} {last} — {d["error"][:50]}')
            rows.append({'first': first, 'last': last, 'error': d['error']})
            continue

        r = d['reason']
        fallback = bool(d.get('meaning_unavailable'))
        short = d.get('meaning_short') or ''
        meaning = d.get('meaning_en') or ''

        meaning_error = d.get('meaning_error') or ''
        src = meaning_source(d.get('given'), meaning, meaning_error,
                             d.get('meaning_source'))

        # 검수(LLM) 결과. app.py 가 meaning_review 로 내려준다.
        rev = d.get('meaning_review') or {}
        rev_status = rev.get('status') or ''
        rev_issues = rev.get('issues') or []
        if isinstance(rev_issues, str):
            rev_issues = [rev_issues]

        # 눈에 띄는 문제만 표시
        flags = []
        # 템플릿이 왜 나왔는지를 플래그에 박아 둔다. 예전에는 '템플릿(LLM실패)'
        # 까지만 남고 원인(meaning_error)은 판정에만 쓰고 버려서, 재시도를
        # 늘려야 할지 타임아웃을 늘려야 할지 알 수 없었다.
        if src == 'template':
            flags.append(f'템플릿(LLM실패:{meaning_error or "원인미상"})')
        if rev_status in ('rejected', 'failed'):
            flags.append(f'검수{rev_status}')
        # 글자별 표기 형식('혜 (惠, hye)')을 지켰는가. LLM 이 새로 쓴 설명만
        # 본다 — 사전 602개는 사람이 쓴 문장이고, 템플릿은 코드가 만들므로
        # 형식이 어긋날 수 없다.
        if src == 'llm':
            bad_form = char_form_misses(d.get('given') or '',
                                        d.get('hanja') or '', meaning)
            if bad_form:
                flags.append(f'표기형식({",".join(bad_form)})')
        # 3인칭 규칙 위반 — 생성 프롬프트와 검수 체크리스트 모두의 금지 항목
        if src == 'llm' and SECOND_PERSON.search(meaning):
            flags.append('2인칭(you)')
        # 프롬프트·한자 뜻이 바뀐 뒤 재생성에 실패해 예전 캐시가 나온 경우.
        # 새로 만든 것과 구분되지 않으면 틀린 뜻이 조용히 계속 나간다.
        if d.get('meaning_stale'):
            flags.append(f"예전캐시({d.get('meaning_stale_why') or '?'})")
        if fallback: flags.append('설명없음')
        # 문법 검사는 템플릿이 조립한 문장에만 적용한다.
        # 사전 602개는 사람이 쓴 문장이라 검사 대상이 아니다.
        if src == 'template' and BAD_EN.search(meaning): flags.append('영어문법')
        if HANGUL.search(short): flags.append('한줄에한국어')
        if HANJA.search(short): flags.append('한줄에한자')
        if LATIN_EXT.search(short) or LATIN_EXT.search(meaning): flags.append('발음기호')
        if not d.get('surname_desc'): flags.append('성씨유래없음')
        if flags: warn += 1

        label = {'llm': 'LLM', 'dict': '사전602', 'native': '순우리말',
                 'template': '템플릿', 'none': '없음'}.get(src, src)
        say(f'\n{i:3d}. {first} {last} ({sex}){"  ⚠ " + ", ".join(flags) if flags else ""}')
        say(f'     음차   : {r["translit"]["hangul"]}')
        say(f'     이름   : {d["full_hangul"]} ({d["full_rom"]})  '
            f'{(d.get("surname_hanja") or "") + (d.get("hanja") or "")}  [{d.get("quality")}]')
        say(f'     한 줄  : {short}')
        # 자르지 않는다. 예전에는 220자에서 잘라 놓고 잘렸다는 표시가 없어,
        # LLM 이 문장을 끝맺지 못한 것처럼 보였다(실제로 그렇게 오해했다).
        say(f'     의미   : [{label}] {meaning}')
        if rev_status:
            say(f'     검수   : {rev_status}'
                + (f' — {"; ".join(str(x) for x in rev_issues)}' if rev_issues else ''))

        rows.append({
            'first': first, 'last': last, 'sex': sex,
            'translit': r['translit']['hangul'],
            'full': d['full_hangul'], 'rom': d['full_rom'],
            'hanja': (d.get('surname_hanja') or '') + (d.get('hanja') or ''),
            'quality': d.get('quality'),
            'source': src,
            # 템플릿으로 떨어진 이유. 이것이 없어 원인을 추적할 수 없었다.
            'meaning_error': meaning_error,
            'review': rev_status,
            'review_issues': ' / '.join(str(x) for x in rev_issues),
            'native': bool(d.get('is_native')),
            'short': short, 'meaning': meaning,
            'surname_desc': d.get('surname_desc') or '',
            'matches': ' / '.join(f'{m["src"]}→{m["tgt"]} {m["style"]["label"]}'
                                  for m in r['matches']),
            'fallback': fallback,
            'flags': flags,
        })

    ok_rows = [r for r in rows if 'error' not in r]
    n = len(ok_rows) or 1
    cnt = {}
    for r in ok_rows:
        cnt[r.get('source')] = cnt.get(r.get('source'), 0) + 1
        cnt[r.get('quality')] = cnt.get(r.get('quality'), 0) + 1

    say('\n' + '=' * 74)
    say(f'완료 {len(rows)}건 · 변환 실패 {len(rows) - len(ok_rows)}건 '
        f'· 확인 필요 {warn}건')
    say('')
    say('의미 설명 출처')
    say(f'  LLM 생성   {cnt.get("llm", 0):4d}건  {100*cnt.get("llm",0)/n:5.1f}%'
        '   <- 이 값이 나와야 LLM 경로가 검증된 것')
    say(f'  사전 602   {cnt.get("dict", 0):4d}건  {100*cnt.get("dict",0)/n:5.1f}%'
        '   (미리 작성된 설명, LLM 검증 아님)')
    say(f'  순우리말   {cnt.get("native", 0):4d}건  {100*cnt.get("native",0)/n:5.1f}%'
        '   (한자가 없어 로컬 생성, 정상)')
    say(f'  템플릿     {cnt.get("template", 0):4d}건  '
        f'{100*cnt.get("template",0)/n:5.1f}%   <- 0 이어야 정상')
    # 템플릿이 나왔으면 원인을 바로 보여준다. 한 줄이라도 유형이 찍혀 있으면
    # 재시도를 늘려야 할지(transient) 타임아웃을 늘려야 할지(timeout) 정할 수 있다.
    tpl = [r for r in ok_rows if r.get('source') == 'template']
    if tpl:
        kinds = {}
        for r in tpl:
            k = r.get('meaning_error') or '원인미상'
            kinds[k] = kinds.get(k, 0) + 1
        say('    실패 원인: ' + ', '.join(f'{k} {v}건' for k, v in
                                       sorted(kinds.items(), key=lambda x: -x[1])))
        say('      transient=과부하·레이트리밋(재시도 소진) · timeout=응답 지연')
        say('      credit=크레딧 소진 · auth=키 문제 · other=요청 오류')
    say('')
    # 검수(LLM) 집계. 검수가 돌았는지 여부가 결과 파일에 남지 않아, 예전에는
    # 캐시를 직접 열어 보지 않으면 적용 여부를 확인할 수 없었다.
    rev_cnt = {}
    for r in ok_rows:
        if r.get('review'):
            rev_cnt[r['review']] = rev_cnt.get(r['review'], 0) + 1
    if rev_cnt:
        say('설명 검수(LLM)')
        for k, ko in (('ok', '문제없음'), ('edited', '수정함'),
                      ('rejected', '수정거부'), ('failed', '검수실패'),
                      ('disabled', '검수꺼짐')):
            if rev_cnt.get(k):
                say(f'  {ko:8s} {rev_cnt[k]:4d}건')
        say('')
    say('이름 품질 등급')
    for q in ('Q1', 'Q2', 'Q3', 'Q4'):
        if cnt.get(q):
            say(f'  {q}         {cnt[q]:4d}건  {100*cnt[q]/n:5.1f}%')
    bad_en = [r for r in ok_rows if '영어문법' in (r.get('flags') or [])]
    if bad_en:
        say('')
        say(f'[!] 영어 문법 의심 {len(bad_en)}건 (관사/품사 오류):')
        for r in bad_en[:5]:
            say(f'    {r["first"]} -> {r["full"]}')
    if not A.TRANSLIT.llm_available:
        say('')
        say('※ API 키 없이 실행했으므로 LLM 생성 0건은 당연합니다.')
        say('  키를 설정하고 다시 돌려야 실제 품질을 확인할 수 있습니다.')
    say('')
    fb = sum(1 for r in rows if r.get('fallback'))
    if fb:
        say(f'폴백 {fb}건 \u2014 API 키가 없거나 호출이 실패한 경우입니다.')

    # 결과 파일 이름 — 국가별로 돌릴 때 서로 덮어쓰지 않도록
    # names_japan.txt 로 돌리면 result_japan.xlsx 로 저장된다.
    if args.out:
        stem = args.out
    elif args.file:
        base = os.path.splitext(os.path.basename(args.file))[0]
        stem = 'result_' + base[6:] if base.startswith('names_') else 'result_' + base
    else:
        stem = 'batch_result'
    def _w_txt(t):
        with open(t, 'w', encoding='utf-8') as f:
            f.write('\n'.join(out_lines))

    def _w_json(t):
        with open(t, 'w', encoding='utf-8') as f:
            json.dump(rows, f, ensure_ascii=False, indent=1)

    saved = [p for p in (_save(stem + '.txt', _w_txt),
                         _save(stem + '.json', _w_json)) if p]

    xlsx = write_xlsx(stem + '.xlsx', rows, cnt, n, warn)
    if xlsx:
        saved.append(xlsx)

    print('\n저장: ' + ' / '.join(saved))
    print('엑셀 파일을 열면 필터·정렬로 검토할 수 있고,')
    print('JSON 내용을 대화에 붙여 넣으면 함께 검토할 수 있습니다.')


# ---------------------------------------------------------------- 엑셀 출력
XL_COLS = [
    ('번호', 6), ('영어 이름', 14), ('성', 14), ('성별', 6),
    ('음차', 12), ('한국 이름', 12), ('로마자', 16), ('한자', 10),
    ('품질', 7), ('의미 출처', 10), ('실패 원인', 12), ('검수', 10),
    ('한 줄 의미', 30),
    ('의미 설명', 70), ('성씨 유래', 40), ('음절 매칭', 26), ('확인 플래그', 24),
    ('검수 지적', 50),
]
# 줄바꿈해서 보여줄 긴 칸. 이름으로 찾는다 — 칸을 추가할 때 번호를 다시
# 세지 않아도 되고, 세다가 틀려 엉뚱한 칸이 줄바꿈되는 일도 없다.
_WRAP_COLS = {'한 줄 의미', '의미 설명', '성씨 유래', '음절 매칭',
              '확인 플래그', '검수 지적'}
_WRAP_IDX = {i for i, (nm, _w) in enumerate(XL_COLS, 1) if nm in _WRAP_COLS}
_SRC_IDX = next(i for i, (nm, _w) in enumerate(XL_COLS, 1) if nm == '의미 출처')
SRC_KO = {'llm': 'LLM', 'dict': '사전602', 'native': '순우리말', 'template': '템플릿',
          'none': '없음', 'error': '오류'}
REV_KO = {'ok': '문제없음', 'edited': '수정함', 'rejected': '수정거부',
          'failed': '검수실패', 'disabled': '검수꺼짐', '': ''}


def write_xlsx(path, rows, cnt, n, warn):
    """결과를 엑셀로. openpyxl 이 없으면 건너뛴다(다른 출력은 그대로 남는다)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        print('\n(openpyxl 이 없어 엑셀 출력을 건너뜁니다. '
              'pip install openpyxl 후 다시 실행하세요.)')
        return None

    base = Font(name='Arial', size=10)
    head_f = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', fgColor='4A4038')
    fills = {'llm': PatternFill('solid', fgColor='D9F2E9'),
             'native': PatternFill('solid', fgColor='EFF3D9'),
             'dict': PatternFill('solid', fgColor='E4EDFA'),
             'template': PatternFill('solid', fgColor='FBE0DD'),
             'error': PatternFill('solid', fgColor='EEEEEE')}
    wrap = Alignment(vertical='top', wrap_text=True)
    top = Alignment(vertical='top')

    wb = Workbook()
    ws = wb.active
    ws.title = '결과'

    for c, (name, width) in enumerate(XL_COLS, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font, cell.fill = head_f, head_fill
        cell.alignment = Alignment(vertical='center', horizontal='center')
        ws.column_dimensions[get_column_letter(c)].width = width

    for i, r in enumerate(rows, 1):
        src = r.get('source') or ('error' if 'error' in r else 'none')
        vals = [
            i, r.get('first', ''), r.get('last', ''), r.get('sex', ''),
            r.get('translit', ''), r.get('full', ''), r.get('rom', ''),
            r.get('hanja', ''), r.get('quality', ''), SRC_KO.get(src, src),
            r.get('meaning_error', ''), REV_KO.get(r.get('review', ''), r.get('review', '')),
            r.get('short', ''),
            r.get('meaning', '') or r.get('error', ''),
            r.get('surname_desc', ''), r.get('matches', ''),
            ', '.join(r.get('flags') or []),
            r.get('review_issues', ''),
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i + 1, column=c, value=v)
            cell.font = base
            cell.alignment = wrap if c in _WRAP_IDX else top
        if src in fills:
            ws.cell(row=i + 1, column=_SRC_IDX).fill = fills[src]

    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(XL_COLS))}{len(rows) + 1}'

    # ---- 요약 시트
    s = wb.create_sheet('요약')
    s.column_dimensions['A'].width = 26
    s.column_dimensions['B'].width = 12
    s.column_dimensions['C'].width = 10
    s.column_dimensions['D'].width = 44

    def put(row, label, value=None, pct=None, note=None, bold=False):
        f = Font(name='Arial', size=10, bold=bold)
        s.cell(row=row, column=1, value=label).font = f
        if value is not None:
            s.cell(row=row, column=2, value=value).font = f
        if pct is not None:
            c = s.cell(row=row, column=3, value=pct)
            c.font, c.number_format = f, '0.0%'
        if note:
            s.cell(row=row, column=4, value=note).font = Font(
                name='Arial', size=9, color='777777')

    ok = len([r for r in rows if 'error' not in r])
    put(1, '배치 테스트 요약', bold=True)
    put(3, '총 건수', len(rows))
    put(4, '변환 성공', ok, ok / (len(rows) or 1))
    put(5, '변환 실패', len(rows) - ok, (len(rows) - ok) / (len(rows) or 1),
        '0 이어야 정상')
    put(6, '확인 필요(플래그)', warn, warn / (len(rows) or 1))

    put(8, '의미 설명 출처', bold=True)
    put(9, 'LLM 생성', cnt.get('llm', 0), cnt.get('llm', 0) / n,
        '이 값이 나와야 LLM 경로가 검증된 것')
    put(10, '사전 602', cnt.get('dict', 0), cnt.get('dict', 0) / n,
        '미리 작성된 설명 — LLM 검증 아님')
    put(11, '템플릿', cnt.get('template', 0), cnt.get('template', 0) / n,
        '0 이어야 정상 (LLM 실패)')

    row = 13
    # 템플릿이 나온 이유. 유형이 없으면 손댈 곳을 정할 수 없다.
    tpl_kinds = {}
    for r in rows:
        if r.get('source') == 'template':
            k = r.get('meaning_error') or '원인미상'
            tpl_kinds[k] = tpl_kinds.get(k, 0) + 1
    if tpl_kinds:
        put(row, '템플릿 실패 원인', bold=True)
        row += 1
        notes = {'transient': '과부하·레이트리밋 — 재시도 소진',
                 'timeout': '응답 지연 — GEN_TIMEOUT 을 올린다',
                 'credit': '크레딧 소진', 'auth': 'API 키 문제',
                 'other': '요청 오류', '원인미상': '기록 없음(구버전)'}
        for k, v in sorted(tpl_kinds.items(), key=lambda x: -x[1]):
            put(row, k, v, v / n, notes.get(k, ''))
            row += 1
        row += 1

    # 검수(LLM) 집계 — 검수가 실제로 돌았는지 여기서 바로 확인된다.
    rev_cnt = {}
    for r in rows:
        if r.get('review'):
            rev_cnt[r['review']] = rev_cnt.get(r['review'], 0) + 1
    if rev_cnt:
        put(row, '설명 검수(LLM)', bold=True)
        row += 1
        rev_note = {'ok': '고칠 것 없음', 'edited': '검수가 고쳐서 내보냄',
                    'rejected': '수정본이 검증을 못 통과 — 원문 유지',
                    'failed': '검수 호출 실패 — 원문 유지',
                    'disabled': '검수기가 꺼져 있음'}
        for k in ('ok', 'edited', 'rejected', 'failed', 'disabled'):
            if rev_cnt.get(k):
                put(row, REV_KO.get(k, k), rev_cnt[k], rev_cnt[k] / n,
                    rev_note.get(k, ''))
                row += 1
        row += 1

    put(row, '이름 품질 등급', bold=True)
    row += 1
    for q in ('Q1', 'Q2', 'Q3', 'Q4'):
        if cnt.get(q):
            put(row, q, cnt[q], cnt[q] / n)
            row += 1

    flagcnt = {}
    for r in rows:
        for fl in (r.get('flags') or []):
            flagcnt[fl] = flagcnt.get(fl, 0) + 1
    if flagcnt:
        row += 1
        put(row, '플래그별 건수', bold=True)
        row += 1
        for fl, k in sorted(flagcnt.items(), key=lambda x: -x[1]):
            put(row, fl, k, k / n)
            row += 1

    uniq = len({r.get('full') for r in rows if r.get('full')})
    row += 1
    put(row, '서로 다른 결과 이름', uniq, uniq / (ok or 1), '다양성 — 높을수록 좋음')

    return _save(path, wb.save)


if __name__ == '__main__':
    main()
